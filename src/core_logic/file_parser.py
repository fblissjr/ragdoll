# src/core_logic/file_parser.py
import os
from pathlib import Path
from datetime import datetime
from typing import List, Dict, Any, Tuple, Optional, Callable

# Third-party library imports
import polars as pl
from pypdf import PdfReader
from docx import Document as DocxDocument
import pypandoc
from ebooklib import epub, ITEM_DOCUMENT
from bs4 import BeautifulSoup
import json as standard_json

# RAGdoll project-specific imports
from .. import ragdoll_utils

# --- Magika Integration ---
MAGIKA_AVAILABLE = False
MAGIKA_INSTANCE = None
ContentTypeLabelEnumFromMagika = None # Placeholder for Magika's enum type

try:
    import magika 
    from magika.types import ContentTypeLabel # Import the enum for type safety if needed
    ContentTypeLabelEnumFromMagika = ContentTypeLabel # Assign to module-level variable
    MAGIKA_INSTANCE = magika.Magika(prediction_mode=magika.PredictionMode.HIGH_CONFIDENCE)
    MAGIKA_AVAILABLE = True
    print("File Parser: Magika library loaded successfully.")
except ImportError:
    print("File Parser: Magika library not found. File typing will rely solely on extensions.")
except Exception as e_magika_init:
    print(f"File Parser: Error initializing Magika: {e_magika_init}. File typing will rely on extensions.")
    MAGIKA_INSTANCE = None 
    MAGIKA_AVAILABLE = False

# --- Helper Function for Tabular Data ---
def _extract_tabular_row_as_document(
    row_data: Dict[str, Any], 
    column_names: List[str], 
    base_metadata_for_row: Dict[str, Any], 
    row_index_val: int
) -> Optional[Tuple[str, Dict[str, Any]]]:
    text_parts = [
        f"{col_name}: {str(row_data.get(col_name))}" 
        for col_name in column_names 
        if row_data.get(col_name) is not None and str(row_data.get(col_name)).strip()
    ]
    content_string = " | ".join(text_parts)
    cleaned_content_string = ragdoll_utils.clean_text(content_string)
    if cleaned_content_string:
        row_specific_metadata = {**base_metadata_for_row, "row_index": row_index_val}
        return cleaned_content_string, row_specific_metadata
    return None

# --- Recursive JSON Value Extractor ---
def _extract_json_value_recursive(
    value: Any, 
    current_json_path: str, 
    path_separator: str = "."
) -> List[Tuple[str, str]]: 
    texts_with_paths: List[Tuple[str, str]] = []
    if isinstance(value, dict):
        if not value: return []
        for key, item_value in value.items():
            new_json_path = f"{current_json_path}{path_separator}{key}" if current_json_path != "root" else key
            texts_with_paths.extend(_extract_json_value_recursive(item_value, new_json_path, path_separator))
    elif isinstance(value, list):
        if not value: return []
        for index, list_item_value in enumerate(value):
            new_json_path = f"{current_json_path}{path_separator}[{index}]" if current_json_path != "root" else f"[{index}]"
            texts_with_paths.extend(_extract_json_value_recursive(list_item_value, new_json_path, path_separator))
    elif isinstance(value, (str, int, float, bool)):
        cleaned_text_value = ragdoll_utils.clean_text(str(value))
        if cleaned_text_value:
            texts_with_paths.append((cleaned_text_value, current_json_path))
    return texts_with_paths

# --- Individual File Type Extraction Functions ---
def _extract_from_txt(filepath: str, base_metadata: Dict[str, Any]) -> List[Tuple[str, Dict[str, Any]]]:
    try:
        with open(filepath, 'r', encoding='utf-8', errors='replace') as f:
            content_str = f.read()
        cleaned_content = ragdoll_utils.clean_text(content_str)
        if cleaned_content:
            return [(cleaned_content, base_metadata.copy())] 
    except Exception as e:
        print(f"  File Parser Error (_extract_from_txt) for {filepath}: {e}")
    return []

def _extract_from_md(filepath: str, base_metadata: Dict[str, Any]) -> List[Tuple[str, Dict[str, Any]]]:
    base_metadata_md = base_metadata.copy()
    base_metadata_md["content_type"] = "markdown"
    return _extract_from_txt(filepath, base_metadata_md)

def _extract_from_code_text(filepath: str, base_metadata: Dict[str, Any]) -> List[Tuple[str, Dict[str, Any]]]:
    base_metadata_code = base_metadata.copy()
    base_metadata_code["content_type"] = "code" 
    return _extract_from_txt(filepath, base_metadata_code)

def _extract_from_html(filepath: str, base_metadata: Dict[str, Any]) -> List[Tuple[str, Dict[str, Any]]]:
    try:
        with open(filepath, 'r', encoding='utf-8', errors='replace') as f:
            html_content = f.read()
        soup = BeautifulSoup(html_content, 'html.parser')
        for script_or_style in soup(["script", "style"]):
            script_or_style.decompose()
        text_content = soup.get_text(separator='\n', strip=True)
        cleaned_content = ragdoll_utils.clean_text(text_content)
        if cleaned_content:
            html_meta = base_metadata.copy()
            html_meta["content_type"] = "html_text"
            return [(cleaned_content, html_meta)]
    except Exception as e:
        print(f"  File Parser Error (_extract_from_html) for {filepath}: {e}")
    return []

def _extract_from_xml(filepath: str, base_metadata: Dict[str, Any]) -> List[Tuple[str, Dict[str, Any]]]:
    try:
        with open(filepath, 'r', encoding='utf-8', errors='replace') as f:
            xml_content = f.read()
        soup = BeautifulSoup(xml_content, 'xml') 
        text_content = soup.get_text(separator='\n', strip=True)
        cleaned_content = ragdoll_utils.clean_text(text_content)
        if cleaned_content:
            xml_meta = base_metadata.copy()
            xml_meta["content_type"] = "xml_text"
            return [(cleaned_content, xml_meta)]
    except Exception as e:
        print(f"  File Parser Error (_extract_from_xml) for {filepath}: {e}")
    return []

def _extract_from_json(filepath: str, base_metadata: Dict[str, Any]) -> List[Tuple[str, Dict[str, Any]]]:
    extracted_parts: List[Tuple[str, Dict[str, Any]]] = []
    try:
        with open(filepath, 'r', encoding='utf-8', errors='replace') as f:
            data = standard_json.load(f)
        texts_with_json_paths = _extract_json_value_recursive(data, current_json_path="root")
        for text_content, json_path_str in texts_with_json_paths:
            part_meta = {**base_metadata, "json_path": json_path_str, "content_type": "json_text_value"}
            extracted_parts.append((text_content, part_meta))
    except standard_json.JSONDecodeError as e_json:
        print(f"  File Parser Error (_extract_from_json): Invalid JSON in {filepath}: {e_json}")
    except Exception as e:
        print(f"  File Parser Error (_extract_from_json) for {filepath}: {e}")
    return extracted_parts

def _extract_from_jsonl(filepath: str, base_metadata: Dict[str, Any]) -> List[Tuple[str, Dict[str, Any]]]:
    extracted_parts: List[Tuple[str, Dict[str, Any]]] = []
    try:
        with open(filepath, 'r', encoding='utf-8', errors='replace') as f:
            for line_num, line_content in enumerate(f, 1):
                if not line_content.strip(): continue
                try:
                    data_item = standard_json.loads(line_content)
                    texts_with_json_paths = _extract_json_value_recursive(data_item, current_json_path="root")
                    for text_content, json_path_str in texts_with_json_paths:
                        part_meta = {**base_metadata, "line_number": line_num, "json_path": json_path_str, "content_type": "jsonl_text_value"}
                        extracted_parts.append((text_content, part_meta))
                except standard_json.JSONDecodeError:
                    print(f"  File Parser Warning: Invalid JSON in line {line_num} of {filepath}. Treating as plain text.")
                    cleaned_line_text = ragdoll_utils.clean_text(line_content)
                    if cleaned_line_text:
                        part_meta = {**base_metadata, "line_number": line_num, "content_type": "jsonl_malformed_line_text"}
                        extracted_parts.append((cleaned_line_text, part_meta))
    except Exception as e:
        print(f"  File Parser Error (_extract_from_jsonl) for {filepath}: {e}")
    return extracted_parts

def _extract_from_csv(filepath: str, base_metadata: Dict[str, Any]) -> List[Tuple[str, Dict[str, Any]]]:
    extracted_parts: List[Tuple[str, Dict[str, Any]]] = []
    try:
        df = pl.read_csv(filepath, infer_schema_length=1000, truncate_ragged_lines=True, ignore_errors=True, encoding="utf-8-lossy")
        column_names = df.columns
        for i, row_dict in enumerate(df.iter_rows(named=True)):
            doc_tuple_from_row = _extract_tabular_row_as_document(row_dict, column_names, base_metadata, i)
            if doc_tuple_from_row: extracted_parts.append(doc_tuple_from_row)
    except Exception as e:
        print(f"  File Parser Error (_extract_from_csv) for {filepath}: {e}")
    return extracted_parts

def _extract_from_tsv(filepath: str, base_metadata: Dict[str, Any]) -> List[Tuple[str, Dict[str, Any]]]:
    extracted_parts: List[Tuple[str, Dict[str, Any]]] = []
    try:
        df = pl.read_csv(filepath, separator='\t', infer_schema_length=1000, truncate_ragged_lines=True, ignore_errors=True, encoding="utf-8-lossy")
        column_names = df.columns
        for i, row_dict in enumerate(df.iter_rows(named=True)):
            doc_tuple_from_row = _extract_tabular_row_as_document(row_dict, column_names, base_metadata, i)
            if doc_tuple_from_row: extracted_parts.append(doc_tuple_from_row)
    except Exception as e:
        print(f"  File Parser Error (_extract_from_tsv) for {filepath}: {e}")
    return extracted_parts

def _extract_from_parquet(filepath: str, base_metadata: Dict[str, Any]) -> List[Tuple[str, Dict[str, Any]]]:
    extracted_parts: List[Tuple[str, Dict[str, Any]]] = []
    try:
        df = pl.read_parquet(filepath)
        column_names = df.columns
        for i, row_dict in enumerate(df.iter_rows(named=True)):
            doc_tuple_from_row = _extract_tabular_row_as_document(row_dict, column_names, base_metadata, i)
            if doc_tuple_from_row: extracted_parts.append(doc_tuple_from_row)
    except Exception as e:
        print(f"  File Parser Error (_extract_from_parquet) for {filepath}: {e}")
    return extracted_parts

def _extract_from_xlsx(filepath: str, base_metadata: Dict[str, Any]) -> List[Tuple[str, Dict[str, Any]]]:
    extracted_parts: List[Tuple[str, Dict[str, Any]]] = []
    try:
        all_sheets_data = pl.read_excel(filepath, sheet_name=None, engine='openpyxl', read_options={"infer_schema_length": 1000})
        for sheet_name_original, df_sheet in all_sheets_data.items():
            if df_sheet.height == 0: continue
            sheet_base_meta = {**base_metadata, "sheet_name": ragdoll_utils.sanitize_filename_for_id(str(sheet_name_original))}
            column_names = df_sheet.columns
            for i, row_dict in enumerate(df_sheet.iter_rows(named=True)):
                doc_tuple_from_row = _extract_tabular_row_as_document(row_dict, column_names, sheet_base_meta, i)
                if doc_tuple_from_row: extracted_parts.append(doc_tuple_from_row)
    except Exception as e_xlsx:
        print(f"  File Parser Error (_extract_from_xlsx) for {filepath}: {e_xlsx}")
    return extracted_parts

def _extract_from_pdf(filepath: str, base_metadata: Dict[str, Any]) -> List[Tuple[str, Dict[str, Any]]]:
    extracted_parts: List[Tuple[str, Dict[str, Any]]] = []
    try:
        reader = PdfReader(filepath)
        if not reader.pages: 
            print(f"  File Parser Info: No pages found in PDF {filepath}")
            return []
        for i, page_obj in enumerate(reader.pages):
            page_text_content = page_obj.extract_text() or ""
            cleaned_page_text = ragdoll_utils.clean_text(page_text_content)
            if cleaned_page_text:
                page_specific_meta = {**base_metadata, "page_number": i + 1, "content_type": "pdf_page_text"}
                extracted_parts.append((cleaned_page_text, page_specific_meta))
    except Exception as e_pdf:
        print(f"  File Parser Error (_extract_from_pdf) processing PDF {filepath}: {e_pdf}")
    return extracted_parts

def _extract_from_docx(filepath: str, base_metadata: Dict[str, Any]) -> List[Tuple[str, Dict[str, Any]]]:
    try:
        markdown_content = pypandoc.convert_file(filepath, 'markdown_strict', format='docx', extra_args=['--wrap=none'])
        cleaned_md_content = ragdoll_utils.clean_text(markdown_content)
        if cleaned_md_content:
            docx_meta = base_metadata.copy()
            docx_meta["content_type"] = "markdown" 
            docx_meta["conversion_source"] = "pandoc_from_docx"
            print(f"  File Parser Info: Successfully converted DOCX '{os.path.basename(filepath)}' to Markdown using Pandoc.")
            return [(cleaned_md_content, docx_meta)]
    except (OSError, RuntimeError, Exception) as e_pandoc: 
        if isinstance(e_pandoc, OSError): # Pandoc not found
            print(f"  File Parser Info: Pandoc not found for DOCX '{os.path.basename(filepath)}': {e_pandoc}. "
                  "Ensure Pandoc is installed and in PATH for optimal DOCX (to Markdown) processing.")
        else: # Other Pandoc errors
             print(f"  File Parser Warning: Pandoc conversion failed for DOCX '{os.path.basename(filepath)}': {e_pandoc}.")
    
    print(f"  Attempting fallback DOCX text extraction for '{os.path.basename(filepath)}' using python-docx...")
    try:
        doc = DocxDocument(filepath)
        full_text_content = "\n".join([p.text for p in doc.paragraphs if p.text and p.text.strip()])
        cleaned_content = ragdoll_utils.clean_text(full_text_content)
        if cleaned_content:
            docx_meta = base_metadata.copy()
            docx_meta["content_type"] = "docx_plain_text"
            docx_meta["conversion_source"] = "python-docx_fallback"
            return [(cleaned_content, docx_meta)]
    except Exception as e_docx:
        print(f"  File Parser Error (_extract_from_docx - python-docx fallback) for {filepath}: {e_docx}")
    return []

def _extract_from_epub(filepath: str, base_metadata: Dict[str, Any]) -> List[Tuple[str, Dict[str, Any]]]:
    extracted_parts: List[Tuple[str, Dict[str, Any]]] = []
    try:
        book = epub.read_epub(filepath)
        for item_obj in book.get_items_of_type(ITEM_DOCUMENT):
            raw_content_bytes = item_obj.get_content()
            try: raw_content_str = raw_content_bytes.decode('utf-8')
            except UnicodeDecodeError: raw_content_str = raw_content_bytes.decode('latin-1', errors='replace')
            soup = BeautifulSoup(raw_content_str, 'html.parser')
            text_content = soup.get_text(separator='\n', strip=True)
            cleaned_content = ragdoll_utils.clean_text(text_content)
            if cleaned_content:
                epub_item_meta = {**base_metadata, "epub_item_id": item_obj.get_id(), "epub_item_name": item_obj.get_name(), "content_type": "epub_xhtml_content"}
                try:
                    title_meta_list = book.get_metadata("DC", "title") 
                    if title_meta_list and title_meta_list[0]: epub_item_meta["epub_book_title"] = title_meta_list[0][0]
                except: pass 
                extracted_parts.append((cleaned_content, epub_item_meta))
    except Exception as e_epub:
        print(f"  File Parser Error (_extract_from_epub) for {filepath}: {e_epub}")
    return extracted_parts

# --- Dispatcher Map for Parsers ---
PARSER_DISPATCH_MAP: Dict[str, Callable[[str, Dict[str, Any]], List[Tuple[str, Dict[str, Any]]]]] = {
    ".txt": _extract_from_txt, ".md": _extract_from_md, ".py": _extract_from_code_text,
    ".js": _extract_from_code_text, ".java": _extract_from_code_text, ".c": _extract_from_code_text,
    ".cpp": _extract_from_code_text, ".cs": _extract_from_code_text, ".go": _extract_from_code_text,
    ".rb": _extract_from_code_text, ".php": _extract_from_code_text, ".html": _extract_from_html,
    ".xml": _extract_from_xml, ".json": _extract_from_json, ".jsonl": _extract_from_jsonl,
    ".csv": _extract_from_csv, ".tsv": _extract_from_tsv, ".parquet": _extract_from_parquet,
    ".xlsx": _extract_from_xlsx, ".pdf": _extract_from_pdf, ".docx": _extract_from_docx,
    ".epub": _extract_from_epub,
}

# --- Main File Type Detection and Extraction Function ---
def _determine_file_type_with_magika(filepath_str: str, extension_type: str) -> Tuple[str, Optional[str], Optional[str], float]:
    if not MAGIKA_AVAILABLE or MAGIKA_INSTANCE is None:
        return extension_type, None, None, 0.0 

    magika_ct_label_str: Optional[str] = None
    magika_group_str: Optional[str] = None
    magika_score_val: float = 0.0
    effective_parser_key = extension_type 

    try:
        file_path_obj = Path(filepath_str)
        # identify_paths returns a list, even for a single path
        results_list = MAGIKA_INSTANCE.identify_paths([file_path_obj])
        
        if not results_list:
            print(f"  File Parser Info: Magika returned no result for {filepath_str}.")
            return extension_type, None, None, 0.0
        
        result: magika.MagikaResult = results_list[0]

        if result.ok and result.output and result.prediction:
            # Correct access based on Magika documentation and tests:
            # result.output is ContentTypeInfo
            # result.output.label is ContentTypeLabel (enum) -> convert to string
            # result.output.group is str
            # result.score is float (shortcut for result.prediction.score)
            
            if result.output.label is not None: # Check if label itself is not None
                 magika_ct_label_str = str(result.output.label) 
            
            magika_score_val = result.score 
            magika_group_str = result.output.group
            
            magika_label_to_ragdoll_parser_key = {
                "markdown": ".md", "python": ".py", "javascript": ".js", "json": ".json",
                "xml": ".xml", "html": ".html", "csv": ".csv", "pdf": ".pdf",
            }
            CONFIDENCE_THRESHOLD_FOR_OVERRIDE = 0.80

            if magika_ct_label_str and \
               magika_ct_label_str in magika_label_to_ragdoll_parser_key and \
               magika_score_val >= CONFIDENCE_THRESHOLD_FOR_OVERRIDE:
                potential_override_type = magika_label_to_ragdoll_parser_key[magika_ct_label_str]
                if potential_override_type in PARSER_DISPATCH_MAP and \
                   (potential_override_type != extension_type or extension_type in [".txt", ".dat", "", None]):
                    effective_parser_key = potential_override_type
                    print(f"  Magika Override: Using parser type '{effective_parser_key}' for '{os.path.basename(filepath_str)}' "
                          f"(original ext: '{extension_type}', Magika label: '{magika_ct_label_str}', score: {magika_score_val:.2f})")
            
            if magika_ct_label_str:
                 print(f"  Magika for {os.path.basename(filepath_str)}: Label='{magika_ct_label_str}' "
                       f"(Group='{magika_group_str}'), Score={magika_score_val:.2f}. "
                       f"Effective parser key chosen: '{effective_parser_key}'.")
            elif result.ok:
                print(f"  Magika for {os.path.basename(filepath_str)}: No specific content label from Magika (e.g., EMPTY/UNKNOWN). "
                       f"Using parser key: '{effective_parser_key}'.")
        else: 
            status_msg = str(result.status) if result and hasattr(result, 'status') else "unknown status/missing output"
            print(f"  File Parser Warning: Magika analysis !ok or missing output for {filepath_str}: {status_msg}. Using extension '{extension_type}'.")
            return extension_type, None, None, 0.0
    except AttributeError as ae: 
        print(f"  File Parser CRITICAL ERROR: Magika result attribute mismatch for {filepath_str}: {ae}. "
              f"Check Magika API. Using extension '{extension_type}'.")
        return extension_type, None, None, 0.0 
    except Exception as e_magika_call:
        print(f"  File Parser Error: General exception during Magika analysis for {filepath_str}: {e_magika_call}. Using extension '{extension_type}'.")
        return extension_type, None, None, 0.0

    return effective_parser_key, magika_ct_label_str, magika_group_str, magika_score_val


def extract_documents_from_file(filepath: str, relative_path: str) -> List[Tuple[str, Dict[str, Any]]]:
    docs_with_metadata: List[Tuple[str, Dict[str, Any]]] = []
    file_path_obj = Path(filepath)
    extension_based_file_type = file_path_obj.suffix.lower()

    effective_parser_key_to_use, magika_label, magika_grp, magika_scr = \
        _determine_file_type_with_magika(filepath, extension_based_file_type)

    base_metadata: Dict[str, Any] = {
        "source_file": relative_path, "file_extension": extension_based_file_type, 
        "detected_content_type_magika": magika_label, "detected_group_magika": magika_grp,           
        "detection_score_magika": round(magika_scr, 4) if magika_scr is not None else None,
        "parsing_type_used": effective_parser_key_to_use, 
        "file_last_modified": datetime.fromtimestamp(file_path_obj.stat().st_mtime).isoformat(),
        "file_size_bytes": file_path_obj.stat().st_size,
    }

    try:
        if not base_metadata["file_size_bytes"] > 0:
            print(f"  File Parser Info: Skipping empty file (0 bytes): {filepath}")
            return []

        if effective_parser_key_to_use in PARSER_DISPATCH_MAP:
            parser_function_to_call = PARSER_DISPATCH_MAP[effective_parser_key_to_use]
            extracted_parts_list = parser_function_to_call(filepath, base_metadata)
            docs_with_metadata.extend(extracted_parts_list)
        else:
            if magika_grp == "text": # Magika thinks it's text, but we don't have a specific parser
                 print(f"  File Parser Info: No specific parser for '{effective_parser_key_to_use}' (ext: '{extension_based_file_type}') for '{filepath}'. "
                       f"Magika identified as '{magika_grp}' group. Attempting generic text extraction.")
                 try:
                     with open(filepath, 'r', encoding='utf-8', errors='replace') as f_generic_text:
                         content_str_generic = f_generic_text.read()
                     cleaned_content_generic = ragdoll_utils.clean_text(content_str_generic)
                     if cleaned_content_generic:
                         generic_meta_text = base_metadata.copy()
                         generic_meta_text["extraction_method"] = "generic_text_fallback_magika_text_group"
                         docs_with_metadata.append((cleaned_content_generic, generic_meta_text))
                 except Exception as e_generic_text:
                     print(f"    Generic text extraction failed for {filepath}: {e_generic_text}")
            else: # Not text by Magika, and no parser by extension/override
                print(f"  File Parser Info: Skipping unsupported file type '{effective_parser_key_to_use}' (original ext: '{extension_based_file_type}', Magika group: '{magika_grp}') for file: {filepath}")
    except Exception as e_main_extraction_loop:
        print(f"  File Parser Error: Major error during extraction dispatch for {filepath} (parser key: {effective_parser_key_to_use}): {e_main_extraction_loop}")
        error_meta_entry = {**base_metadata, "parsing_error_main_dispatch": str(e_main_extraction_loop)}
        docs_with_metadata.append(("", error_meta_entry))
    
    return docs_with_metadata